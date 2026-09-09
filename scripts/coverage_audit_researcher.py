#!/usr/bin/env python3
"""Did every LNG unit actually get updated this cycle, and by whom?

Answers the pre-handoff question "is the 2026 update complete?" against three
sources that disagree with each other, which is the whole point:

  1. **The assignment sheet** ("GGIT LNG 2026 Research Tracking" ->
     `Researcher Assignments`) — who was *responsible* for each country. Read
     live over the read-only `gws-gem` profile; cached to
     `docs/reference/researcher_assignments.json` so a run works offline.
     The export's own `Researcher` column only says who last *touched* a row,
     which is a different question and routinely a different person.
  2. **`unit_update`** (read-only Postgres) — the real per-unit tick log behind
     the export's `LastUpdated`/`Researcher` columns: every time a researcher
     marked a unit, with the date, the person, and the `research_status`
     (`added` / `updated` / `no changes` / `in progress` / ...). The export
     keeps only the latest tick; this keeps all of them, so "checked and found
     nothing to change" (`no changes`) is visible as the completed work it is
     rather than reading as a gap.
  3. **`plant_history`** — the field-level snapshot log, which is **plant-level**:
     its `plantJSON` embeds the unit array, so unit edits are visible, but a
     revision can't be attributed to one unit of a multi-unit terminal. Read
     `field_edits_since` as "something on this terminal changed", not "this
     unit changed". A researcher can edit
     fields without ticking the box (coordinate passes do this constantly, see
     `location_edit_provenance.py`) and can tick the box without editing
     anything. Comparing 2 against 3 separates those two failure modes instead
     of lumping them into one "stale" count.

Staleness definition (`--since`, default 2026-01-01): a unit is flagged when it
has no tick on or after that date. "Never touched in 2026" is the default
because a fixed trailing window (e.g. 4 months) misclassifies finished work —
a dense pass that wrapped up in March is complete, not stale. `--window-days`
adds the trailing-window count alongside it for comparison, never in place.

    python coverage_audit_researcher.py                       # memo to stdout
    python coverage_audit_researcher.py --xlsx ../batches/coverage_audit.xlsx
    python coverage_audit_researcher.py --memo ../batches/coverage_<stamp>.md
    python coverage_audit_researcher.py --window-days 120     # add 4-month view
    python coverage_audit_researcher.py --offline             # cached roster
    python coverage_audit_researcher.py --researcher Andrew   # scope the tables

Read-only throughout: no live-DB writes, no Sheet writes (the `gws-gem`
profile has read-only scopes by construction).
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import paths  # noqa: E402

ROSTER_SHEET_ID = "1NmdDl3NzoRhu1ZgasYNf5nGr79j4S7aI5mH_utJYdPE"
ROSTER_TAB = "Researcher Assignments"
ROSTER_RANGE = f"'{ROSTER_TAB}'!A1:G200"
ROSTER_CACHE = paths.repo_root() / "docs" / "reference" / "researcher_assignments.json"
GWS_CONFIG_DIR = "~/.config/gws-gem"          # READ-ONLY scopes — never gws-gem-write

DEFAULT_SINCE = "2026-01-01"

# Scope gate. The export's `Fuel` column carries four non-LNG values that share
# the table but are not this tracker's LNG work:
#   Oil  (15 rows / 14 US terminals) — legacy crude & NGL deepwater-port records.
#   NH3  (6)  / LH2 (2) / eLNG (1)   — alt-fuel unit rows on real European LNG
#        terminals, whose fields (NH3, LH2, SyntheticLNG) the methodology marks
#        "no longer updated as of 2026", i.e. read-only.
# Flagging any of them as "not updated" would manufacture a gap out of work
# nobody is supposed to be doing. Excluded by default, counted explicitly in the
# memo so the exclusion is visible rather than silent.
IN_SCOPE_FUELS = {"lng"}

# `research_status` rows that mean a researcher looked at the unit and finished
# with it. `no changes` is completed work: the check happened, nothing needed
# editing. `in progress` is explicitly NOT done.
STATUS_DONE = {"added", "updated", "no changes"}

# A never-updated `cancelled` record is a different animal from a never-updated
# `operating` one: dead projects are legitimately deprioritized, and lumping
# them together turns 16 real gaps into an 80-row alarm. Split, never summed.
DEAD_STATUSES = {"cancelled", "shelved", "retired", "mothballed"}

# plantJSON keys that move on every save regardless of whether any researched
# value changed — excluded so a bare re-save doesn't read as a real edit.
VOLATILE_JSON_KEYS = {"updates", "modified", "lastUpdated", "researcher"}

# The assignment sheet's Region column is corrupted (Nigeria -> Europe,
# Colombia -> Africa, Republic of the Congo -> Americas). Subregion and Country
# are reliable; Region is carried through for reference and never joined on.
ROSTER_REGION_UNRELIABLE = True


# ---------------------------------------------------------------- roster


def _gws_read_roster() -> list[list[str]]:
    """Live read of the assignment tab over the read-only work profile."""
    env = dict(os.environ)
    env["GOOGLE_WORKSPACE_CLI_CONFIG_DIR"] = os.path.expanduser(GWS_CONFIG_DIR)
    env["GOOGLE_WORKSPACE_CLI_KEYRING_BACKEND"] = "file"
    proc = subprocess.run(
        ["gws", "sheets", "+read", "--spreadsheet", ROSTER_SHEET_ID,
         "--range", ROSTER_RANGE, "--format", "json"],
        capture_output=True, text=True, env=env, timeout=120,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"gws sheets read failed: {proc.stderr.strip()[:400]}")
    # gws prints a "Using keyring backend: file" banner ahead of the JSON body
    body = proc.stdout[proc.stdout.index("{"):]
    return json.loads(body).get("values", [])


def load_roster(offline: bool) -> tuple[dict, dict, str]:
    """country -> {researcher, subregion, region, projects, complete}.

    Returns (by_country, raw_meta, provenance) — provenance says whether the
    numbers came off the live Sheet or the cache, because a stale roster
    silently reassigns blame.
    """
    values, provenance = None, ""
    if not offline:
        try:
            values = _gws_read_roster()
            provenance = "live Sheet via gws-gem (read-only)"
        except Exception as exc:                                   # noqa: BLE001
            print(f"WARNING: live roster read failed ({exc}); falling back to cache",
                  file=sys.stderr)
    if values is None:
        if not ROSTER_CACHE.exists():
            sys.exit(f"no cached roster at {ROSTER_CACHE} and live read unavailable")
        cached = json.loads(ROSTER_CACHE.read_text())
        values = cached["values"]
        provenance = f"cache {ROSTER_CACHE.name} (fetched {cached.get('fetched','?')})"
    else:
        ROSTER_CACHE.parent.mkdir(parents=True, exist_ok=True)
        ROSTER_CACHE.write_text(json.dumps(
            {"fetched": date.today().isoformat(), "sheet_id": ROSTER_SHEET_ID,
             "tab": ROSTER_TAB, "values": values}, indent=2) + "\n")

    header = [h.strip() for h in values[0]]
    idx = {name: header.index(name) for name in header}

    def cell(row, name):
        i = idx.get(name)
        return (row[i].strip() if i is not None and i < len(row) else "")

    by_country = {}
    for row in values[1:]:
        country = cell(row, "Country")
        if not country:
            continue
        by_country[country] = {
            "researcher": cell(row, "Researcher"),
            "subregion": cell(row, "Subregion"),
            "region": cell(row, "Region"),
            "projects_est": cell(row, "Projects*"),
            "complete_flag": cell(row, "Complete?"),
        }
    return by_country, {"rows": len(by_country)}, provenance


# ---------------------------------------------------------------- database


def _numeric_id(value: str) -> int | None:
    """Export ids are prefixed: T100000130368 / G100002046401."""
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    return int(digits) if digits else None


def fetch_ticks_and_edits():
    """(ticks_by_unit, edits_by_plant, users) straight from the read-only DB."""
    from sqlalchemy import text

    eng = paths.get_engine()
    with eng.connect() as conn:
        users = {
            r[0]: (f"{r[1] or ''} {r[2] or ''}".strip() or r[3])
            for r in conn.execute(text(
                "select id, first_name, last_name, username from auth_user"))
        }
        statuses = {r[0]: r[1] for r in conn.execute(text(
            "select id, option from research_status"))}

        ticks = defaultdict(list)
        for unit_id, when, updater, status_id in conn.execute(text("""
            select uu.unit_id, uu."lastUpdated", uu.updater_id, uu."researchStatus_id"
              from unit_update uu
              join powerplant_unit u on u.id = uu.unit_id
              join plant p on p.id = u.plant_id
             where p."projectType" = 8
             order by uu.unit_id, uu."lastUpdated"
        """)):
            ticks[unit_id].append({
                "date": when.isoformat() if when else "",
                "who": users.get(updater, f"user:{updater}"),
                "status": statuses.get(status_id, str(status_id)),
            })

        # Field-level snapshots: did anything researched actually change?
        snapshots = defaultdict(list)
        for plant_id, editor, modified, blob in conn.execute(text("""
            select h.plant_id, h.editor_id, h.modified, h."plantJSON"
              from plant_history h
              join plant p on p.id = h.plant_id
             where p."projectType" = 8 and h."plantJSON" is not null
             order by h.plant_id, h.modified
        """)):
            snapshots[plant_id].append((editor, modified, blob))

    edits = {}
    for plant_id, revisions in snapshots.items():
        previous, real = None, []
        for editor, modified, blob in revisions:
            payload = blob if isinstance(blob, dict) else json.loads(blob)
            comparable = {k: v for k, v in payload.items()
                          if k not in VOLATILE_JSON_KEYS}
            if previous is not None and comparable != previous:
                real.append({"date": str(modified)[:10],
                             "who": users.get(editor, f"user:{editor}")})
            previous = comparable
        edits[plant_id] = real
    return ticks, edits, users


# ---------------------------------------------------------------- analysis


def split_scope(rows, fuels):
    """(in_scope, out_of_scope_by_fuel) — see IN_SCOPE_FUELS."""
    keep, drop = [], defaultdict(list)
    for row in rows:
        fuel = (row.get("Fuel") or "").strip()
        if fuel.lower() in fuels:
            keep.append(row)
        else:
            drop[fuel or "(blank)"].append(row)
    return keep, drop


def analyse(rows, roster, ticks, edits, since: date, window_days: int | None):
    """One verdict per unit row, plus the two rollups the memo prints."""
    window_start = (date.today() - timedelta(days=window_days)) if window_days else None
    units = []

    for row in rows:
        country = row.get("Country/Area", "").strip()
        assigned = roster.get(country, {})
        unit_id = _numeric_id(row.get("UnitID"))
        plant_id = _numeric_id(row.get("TerminalID"))

        unit_ticks = ticks.get(unit_id, [])
        done_ticks = [t for t in unit_ticks if t["status"] in STATUS_DONE]
        in_since = [t for t in done_ticks if t["date"] and t["date"] >= since.isoformat()]
        in_window = ([t for t in done_ticks
                      if t["date"] and t["date"] >= window_start.isoformat()]
                     if window_start else [])

        plant_edits = edits.get(plant_id, [])
        edits_since = [e for e in plant_edits if e["date"] >= since.isoformat()]

        stale = not in_since
        if stale and edits_since:
            # Fields moved but nobody ticked the box: work happened, the
            # tracking sheet just can't see it.
            verdict = "edited_not_ticked"
        elif stale:
            verdict = "stale"
        elif not edits_since and all(t["status"] != "no changes" for t in in_since):
            # Ticked `added`/`updated` yet no researched field changed —
            # either a no-op save or the tick overstates the work.
            verdict = "ticked_no_edit"
        else:
            verdict = "current"

        units.append({
            "terminal_id": row.get("TerminalID", ""),
            "unit_id": row.get("UnitID", ""),
            "terminal": row.get("TerminalName", ""),
            "unit": row.get("UnitName", ""),
            "country": country,
            "status": row.get("Status", ""),
            "substatus": row.get("Substatus", ""),
            "assigned": assigned.get("researcher", ""),
            "subregion": assigned.get("subregion", ""),
            "complete_flag": assigned.get("complete_flag", ""),
            "export_last_updated": row.get("LastUpdated", "").strip(),
            "export_researcher": row.get("Researcher", "").strip(),
            "last_tick": done_ticks[-1]["date"] if done_ticks else "",
            "last_tick_by": done_ticks[-1]["who"] if done_ticks else "",
            "last_tick_status": done_ticks[-1]["status"] if done_ticks else "",
            "ticks_since": len(in_since),
            "ticks_in_window": len(in_window),
            "last_field_edit": plant_edits[-1]["date"] if plant_edits else "",
            "last_field_edit_by": plant_edits[-1]["who"] if plant_edits else "",
            "field_edits_since": len(edits_since),
            "flagged": stale,
            "verdict": verdict,
            "live_status": row.get("Status", "").strip().lower() not in DEAD_STATUSES,
        })

    return units, window_start


def rollup(units, key):
    out = defaultdict(lambda: {
        "unit_rows": 0, "terminals": set(), "flagged": 0, "in_window_missing": 0,
        "stale": 0, "edited_not_ticked": 0, "ticked_no_edit": 0,
        "stale_live": 0, "stale_dead": 0,
        "oldest": "", "newest": "", "countries": set(),
    })
    for u in units:
        b = out[u[key] or "(unassigned)"]
        b["unit_rows"] += 1
        b["terminals"].add(u["terminal_id"])
        b["countries"].add(u["country"])
        b["flagged"] += int(u["flagged"])
        b["in_window_missing"] += int(u["ticks_in_window"] == 0)
        if u["verdict"] in ("stale", "edited_not_ticked", "ticked_no_edit"):
            b[u["verdict"]] += 1
        if u["verdict"] == "stale":
            b["stale_live" if u["live_status"] else "stale_dead"] += 1
        tick = u["last_tick"]
        if tick:
            b["oldest"] = min(b["oldest"] or tick, tick)
            b["newest"] = max(b["newest"], tick)
    return out


# ---------------------------------------------------------------- output


def write_memo(fh, units, roster, since, window_start, provenance, export_path,
               out_of_scope=None):
    by_country = rollup(units, "country")
    by_researcher = rollup(units, "assigned")
    flagged = [u for u in units if u["flagged"]]
    terminals = {u["terminal_id"] for u in units}

    def w(line=""):
        print(line, file=fh)

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    w(f"# LNG coverage audit — researcher assignment vs actual updates")
    w()
    w(f"Generated {stamp} local. Export `{Path(export_path).name}` "
      f"({len(units)} unit rows, {len(terminals)} terminals, "
      f"{len({u['country'] for u in units})} countries).")
    w(f"Roster source: {provenance}.")
    w(f"Staleness definition: **no `added`/`updated`/`no changes` tick on or after "
      f"{since.isoformat()}** (\"never touched in 2026\").")
    if window_start:
        w(f"Trailing-window comparison also shown: nothing since "
          f"{window_start.isoformat()}.")
    w()
    if out_of_scope:
        parts = ", ".join(f"{fuel} {len(rs)} rows/"
                          f"{len({r['TerminalID'] for r in rs})} terminals"
                          for fuel, rs in sorted(out_of_scope.items()))
        w(f"Scope: `Fuel = LNG` only. Excluded as out of scope ({parts}) — the "
          "Oil rows are legacy US crude/NGL deepwater-port records, and the "
          "NH3/LH2/eLNG rows are alt-fuel units on European LNG terminals whose "
          "fields the methodology marks read-only from 2026. Neither is work a "
          "researcher was assigned this cycle.")
        w()
    w("Three sources, deliberately cross-checked: the assignment sheet (who was "
      "responsible), `unit_update` (every tick, with person and research status), "
      "and `plant_history` (whether any researched field actually changed). The "
      "export's `Researcher`/`LastUpdated` columns answer neither question on "
      "their own.")
    w()

    w("## Headline")
    w()
    counts = defaultdict(int)
    for u in units:
        counts[u["verdict"]] += 1
    stale_live = [u for u in units if u["verdict"] == "stale" and u["live_status"]]
    stale_dead = [u for u in units if u["verdict"] == "stale" and not u["live_status"]]
    w(f"- **{len(stale_live)}** unit rows on a **live status** (operating, "
      f"construction, proposed, idle) never ticked in {since.year} and with no "
      f"field edit either — **this is the actionable gap.**")
    w(f"- **{len(stale_dead)}** more never ticked, but on a dead status "
      f"(cancelled/shelved/retired/mothballed) — deprioritized by design, not a "
      f"coverage failure. Listed separately so the two never get summed.")
    w(f"- **{counts['edited_not_ticked']}** edited in {since.year} but never ticked — "
      f"work done, tracking sheet blind to it.")
    w(f"- **{counts['ticked_no_edit']}** ticked `added`/`updated` in {since.year} with "
      f"no researched field change — verify the tick reflects real work.")
    w(f"- **{counts['current']}** clean.")
    if window_start:
        w(f"- For comparison, **{sum(1 for u in units if u['ticks_in_window'] == 0)}** "
          f"rows have no tick in the trailing "
          f"{(date.today() - window_start).days}-day window — this number is "
          f"inflated by finished passes that simply ended earlier in the year.")
    w()

    w("## By researcher (assignment sheet)")
    w()
    w("Reading the columns: **live gap** is the number that matters. **dead-status "
      "gap** is cancelled/shelved/retired rows nobody was expected to revisit. "
      "**edited, not ticked** means the fields moved in "
      f"{since.year} but the box was never checked — work done, invisible to the "
      "tracking sheet. The three are disjoint and together make up every row "
      f"with no {since.year} tick.")
    w()
    w("| Researcher | Countries | Terminals | Unit rows | **live gap** | "
      "dead-status gap | edited, not ticked | Tick range |")
    w("|---|---|---|---|---|---|---|---|")
    for name, b in sorted(by_researcher.items(),
                          key=lambda kv: (-kv[1]["stale_live"], -kv[1]["flagged"],
                                          kv[0])):
        rng = f"{b['oldest']} → {b['newest']}" if b["oldest"] else "—"
        w(f"| {name} | {len(b['countries'])} | {len(b['terminals'])} | "
          f"{b['unit_rows']} | **{b['stale_live']}** | {b['stale_dead']} | "
          f"{b['edited_not_ticked']} | {rng} |")
    w()

    w("## By country — countries with any flagged row")
    w()
    w("| Country | Assigned | Sheet says complete | Terminals | Unit rows | "
      "**live gap** | dead-status gap | edited, not ticked | Tick range |")
    w("|---|---|---|---|---|---|---|---|---|")
    for country, b in sorted(by_country.items(),
                             key=lambda kv: (-kv[1]["stale_live"],
                                             -kv[1]["flagged"], kv[0])):
        if not b["flagged"]:
            continue
        meta = roster.get(country, {})
        rng = f"{b['oldest']} → {b['newest']}" if b["oldest"] else "—"
        w(f"| {country} | {meta.get('researcher','(not on sheet)')} | "
          f"{meta.get('complete_flag','—')} | {len(b['terminals'])} | "
          f"{b['unit_rows']} | **{b['stale_live']}** | {b['stale_dead']} | "
          f"{b['edited_not_ticked']} | {rng} |")
    w()

    clean = [c for c, b in by_country.items() if not b["flagged"]]
    w(f"{len(clean)} countries are fully current: {', '.join(sorted(clean)) or '—'}")
    w()

    unmatched_export = sorted({u["country"] for u in units if not u["assigned"]})
    export_countries = {u["country"] for u in units}
    unmatched_roster = sorted(c for c in roster if c not in export_countries)
    w("## Roster/export reconciliation")
    w()
    w(f"- Export countries with no row on the assignment sheet "
      f"({len(unmatched_export)}): {', '.join(unmatched_export) or 'none'}")
    w(f"- Assignment-sheet countries with no terminal in the export "
      f"({len(unmatched_roster)}): {', '.join(unmatched_roster) or 'none'}")
    if ROSTER_REGION_UNRELIABLE:
        w("- The sheet's **Region** column is corrupted (Nigeria→Europe, "
          "Colombia→Africa, Republic of the Congo→Americas). Subregion and "
          "Country are reliable and are what this audit joins on.")
    w()

    w("## The live-status gap, row by row")
    w()
    w("Every unit row on a live status with no "
      f"{since.year} tick and no {since.year} field edit. This is the list to "
      "clear before handoff.")
    w()
    w("| Country | Assigned | Terminal | Unit | Status | Last tick | By | "
      "Last field edit | By |")
    w("|---|---|---|---|---|---|---|---|---|")
    for u in sorted(stale_live, key=lambda u: (u["assigned"], u["country"],
                                               u["terminal"], u["unit"])):
        w(f"| {u['country']} | {u['assigned'] or '—'} | {u['terminal']} | "
          f"{u['unit']} | {u['status']} | {u['last_tick'] or 'never'} | "
          f"{u['last_tick_by'] or '—'} | {u['last_field_edit'] or 'never'} | "
          f"{u['last_field_edit_by'] or '—'} |")
    w()

    w("## Every flagged unit row")
    w()
    w("| Country | Assigned | Terminal | Unit | Status | Last tick | By | "
      "Last field edit | By | Verdict |")
    w("|---|---|---|---|---|---|---|---|---|---|")
    for u in sorted(flagged, key=lambda u: (u["assigned"], u["country"],
                                            u["terminal"], u["unit"])):
        w(f"| {u['country']} | {u['assigned'] or '—'} | {u['terminal']} | "
          f"{u['unit']} | {u['status']} | {u['last_tick'] or 'never'} | "
          f"{u['last_tick_by'] or '—'} | {u['last_field_edit'] or 'never'} | "
          f"{u['last_field_edit_by'] or '—'} | {u['verdict']} |")
    w()


def write_xlsx(path, units, roster, since):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    HEADER = PatternFill("solid", fgColor="EEEEEE")
    RED = PatternFill("solid", fgColor="FFE5E5")
    YELLOW = PatternFill("solid", fgColor="FFF8E1")
    BLUE = PatternFill("solid", fgColor="E5F0FF")
    FILLS = {"stale": RED, "edited_not_ticked": YELLOW, "ticked_no_edit": BLUE}

    wb = openpyxl.Workbook()

    def sheet(title, header, records, fill_key=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER
        for rec in records:
            ws.append(rec["row"])
            if fill_key:
                fill = FILLS.get(rec.get(fill_key))
                if fill:
                    for cell in ws[ws.max_row]:
                        cell.fill = fill
        ws.freeze_panes = "A2"
        for col in range(1, len(header) + 1):
            width = max([len(str(header[col - 1]))] +
                        [len(str(r["row"][col - 1])) for r in records] or [10])
            ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 52)
        return ws

    by_researcher = rollup(units, "assigned")
    sheet("by_researcher",
          ["researcher", "countries", "terminals", "unit_rows",
           "live_status_gap", "dead_status_gap", "edited_not_ticked",
           f"never_ticked_{since.year}_total", "ticked_no_edit",
           "oldest_tick", "newest_tick"],
          [{"row": [name, len(b["countries"]), len(b["terminals"]), b["unit_rows"],
                    b["stale_live"], b["stale_dead"], b["edited_not_ticked"],
                    b["flagged"], b["ticked_no_edit"], b["oldest"], b["newest"]]}
           for name, b in sorted(by_researcher.items(),
                                 key=lambda kv: (-kv[1]["stale_live"],
                                                 -kv[1]["flagged"], kv[0]))])

    by_country = rollup(units, "country")
    sheet("by_country",
          ["country", "assigned_researcher", "sheet_complete", "subregion",
           "terminals", "unit_rows", "live_status_gap", "dead_status_gap",
           "edited_not_ticked", f"never_ticked_{since.year}_total",
           "ticked_no_edit", "oldest_tick", "newest_tick"],
          [{"row": [country, roster.get(country, {}).get("researcher", ""),
                    roster.get(country, {}).get("complete_flag", ""),
                    roster.get(country, {}).get("subregion", ""),
                    len(b["terminals"]), b["unit_rows"], b["stale_live"],
                    b["stale_dead"], b["edited_not_ticked"], b["flagged"],
                    b["ticked_no_edit"], b["oldest"], b["newest"]]}
           for country, b in sorted(by_country.items(),
                                    key=lambda kv: (-kv[1]["flagged"], kv[0]))])

    cols = ["country", "assigned", "subregion", "terminal_id", "terminal",
            "unit_id", "unit", "status", "live_status", "substatus",
            "last_tick", "last_tick_by",
            "last_tick_status", "ticks_since", "last_field_edit",
            "last_field_edit_by", "field_edits_since", "export_last_updated",
            "export_researcher", "verdict"]
    sheet("live_status_gap", cols,
          [{"row": [u[c] for c in cols], "verdict": u["verdict"]}
           for u in sorted((u for u in units
                            if u["verdict"] == "stale" and u["live_status"]),
                           key=lambda u: (u["assigned"], u["country"],
                                          u["terminal"]))],
          fill_key="verdict")
    sheet("flagged_units", cols,
          [{"row": [u[c] for c in cols], "verdict": u["verdict"]}
           for u in sorted((u for u in units if u["flagged"]),
                           key=lambda u: (u["assigned"], u["country"], u["terminal"]))],
          fill_key="verdict")
    sheet("all_units", cols,
          [{"row": [u[c] for c in cols], "verdict": u["verdict"]}
           for u in sorted(units, key=lambda u: (u["country"], u["terminal"],
                                                 u["unit"]))],
          fill_key="verdict")

    del wb["Sheet"]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- main


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", default=str(paths.gem_export_csv()),
                    help="fresh all-fields LNG export")
    ap.add_argument("--since", default=DEFAULT_SINCE,
                    help="staleness cutoff (default %(default)s)")
    ap.add_argument("--window-days", type=int, default=None,
                    help="also report a trailing-window count, for comparison")
    ap.add_argument("--researcher", help="limit tables to one assigned researcher")
    ap.add_argument("--country", help="limit tables to one country")
    ap.add_argument("--fuel", default="LNG",
                    help="comma-separated in-scope Fuel values, or ALL "
                         "(default %(default)s)")
    ap.add_argument("--offline", action="store_true",
                    help="use the cached roster instead of reading the Sheet")
    ap.add_argument("--memo", help="write the markdown memo here (default stdout)")
    ap.add_argument("--xlsx", help="also write a workbook here")
    ap.add_argument("--json", help="dump the per-unit verdicts here")
    args = ap.parse_args()

    since = datetime.strptime(args.since, "%Y-%m-%d").date()
    roster, _, provenance = load_roster(args.offline)

    with open(args.csv, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"no rows in {args.csv}")

    if args.fuel.upper() == "ALL":
        in_scope, out_of_scope = rows, {}
    else:
        fuels = {f.strip().lower() for f in args.fuel.split(",") if f.strip()}
        in_scope, out_of_scope = split_scope(rows, fuels)
        if not in_scope:
            sys.exit(f"no rows with Fuel in {sorted(fuels)}")

    ticks, edits, _ = fetch_ticks_and_edits()
    units, window_start = analyse(in_scope, roster, ticks, edits, since,
                                  args.window_days)

    if args.researcher:
        units = [u for u in units
                 if u["assigned"].lower() == args.researcher.lower()]
    if args.country:
        units = [u for u in units if u["country"].lower() == args.country.lower()]
    if not units:
        sys.exit("no unit rows after filtering")

    if args.json:
        Path(args.json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.json).write_text(json.dumps(units, indent=2) + "\n")
    if args.xlsx:
        write_xlsx(args.xlsx, units, roster, since)
        print(f"wrote {args.xlsx}", file=sys.stderr)

    if args.memo:
        Path(args.memo).parent.mkdir(parents=True, exist_ok=True)
        with open(args.memo, "w", encoding="utf-8") as fh:
            write_memo(fh, units, roster, since, window_start, provenance,
                       args.csv, out_of_scope)
        print(f"wrote {args.memo}", file=sys.stderr)
    else:
        write_memo(sys.stdout, units, roster, since, window_start, provenance,
                   args.csv, out_of_scope)


if __name__ == "__main__":
    main()
