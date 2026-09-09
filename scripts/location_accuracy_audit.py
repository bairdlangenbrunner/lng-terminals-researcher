#!/usr/bin/env python3
"""Audit LNG terminal coordinate accuracy against edit recency.

Supports the location-accuracy improvement pass (georeference workflow,
docs/workflows.md §10): which terminals still carry `Accuracy = approximate`,
split by whether someone has touched the record recently.

    python location_accuracy_audit.py                       # gem_export.csv -> xlsx+csv
    python location_accuracy_audit.py --recent-since 2025    # widen "recent"
    python location_accuracy_audit.py --out-dir ../batches/deliverables

Rolls the unit-row export up to one row per terminal: coordinates, Location,
and Accuracy are project-level fields duplicated across unit rows, while
LastUpdated is per-unit (we take the newest, and flag disagreement).

Read-only: touches nothing but the export CSV.
"""

import argparse
import csv
import os
import subprocess
import sys
from collections import defaultdict

PROJECT_URL = "https://gem-project-db.herokuapp.com/projects/edit/{}"

COLS = [
    "country", "terminal_name", "terminal_id", "project_url",
    "accuracy", "accuracy_mixed", "latitude", "longitude", "coord_dp",
    "location", "state_province", "location_ref",
    "last_updated", "last_updated_mixed", "researcher",
    "fuel", "n_units", "statuses", "facility_types", "unit_names", "wiki",
]

WIDTHS = {
    "country": 16, "terminal_name": 34, "terminal_id": 15, "project_url": 46,
    "accuracy": 12, "accuracy_mixed": 9, "latitude": 11, "longitude": 11,
    "coord_dp": 8, "location": 34, "state_province": 18, "location_ref": 40,
    "last_updated": 12, "last_updated_mixed": 9, "researcher": 18,
    "fuel": 10, "n_units": 7, "statuses": 20, "facility_types": 14,
    "unit_names": 30, "wiki": 40,
}

# legacy non-LNG records (crude/NGL deepwater ports) live in the same table but
# are out of LNG scope — same class the missing-year ref-sweep set aside.
LNG_FUELS = {"", "lng", "gas", "natural gas"}


def _dp(value):
    """Decimal places in a coordinate string — a crude precision signal.

    Trailing zeros are padding, not precision: the export writes 7dp for every
    coordinate, so "-93.0100000" is a 2dp (~1 km) value and must grade as one.
    """
    value = (value or "").strip()
    if "." not in value:
        return 0 if value else ""
    return len(value.rsplit(".", 1)[1].rstrip().rstrip("0"))


def _uniq(values):
    out = []
    for v in values:
        v = (v or "").strip()
        if v and v not in out:
            out.append(v)
    return out


def rollup(csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"no rows in {csv_path}")

    by_terminal = defaultdict(list)
    for r in rows:
        by_terminal[r["TerminalID"]].append(r)

    out = []
    for tid, group in by_terminal.items():
        accs = _uniq(r["Accuracy"] for r in group)
        lus = _uniq(r["LastUpdated"] for r in group)
        first = group[0]
        lat, lon = first["Latitude"].strip(), first["Longitude"].strip()
        dps = [_dp(lat), _dp(lon)]
        out.append({
            "country": first["Country/Area"],
            "terminal_name": first["TerminalName"],
            "terminal_id": tid,
            "project_url": PROJECT_URL.format(tid.lstrip("T")),
            # a mixed-accuracy project is a data bug in itself; show the coarsest
            "accuracy": "approximate" if "approximate" in accs else (accs[0] if accs else ""),
            "accuracy_mixed": "MIXED" if len(accs) > 1 else "",
            "latitude": lat,
            "longitude": lon,
            "coord_dp": min(d for d in dps if d != "") if any(d != "" for d in dps) else "",
            "location": first["Location"],
            "state_province": first["State/Province"],
            "location_ref": first["Location [ref]"],
            "last_updated": max(lus) if lus else "",
            "last_updated_mixed": "MIXED" if len(lus) > 1 else "",
            "researcher": "; ".join(_uniq(r["Researcher"] for r in group)),
            "fuel": "; ".join(_uniq(r["Fuel"] for r in group)),
            "n_units": len(group),
            "statuses": "; ".join(_uniq(r["Status"] for r in group)),
            "facility_types": "; ".join(_uniq(r["FacilityType"] for r in group)),
            "unit_names": "; ".join(_uniq(r["UnitName"] for r in group))[:250],
            "wiki": first["Wiki"],
        })
    return len(rows), out


def split(terminals, recent_since):
    approx = [t for t in terminals if t["accuracy"] == "approximate"]
    blank = [t for t in terminals if not t["accuracy"]]

    def year(t):
        try:
            return int(t["last_updated"][:4])
        except ValueError:
            return 0

    # out-of-scope fuels sort to the bottom of each sheet rather than being dropped
    recent = sorted((t for t in approx if year(t) >= recent_since),
                    key=lambda t: (is_offscope(t), _neg(t["last_updated"]), t["country"]))
    stale = sorted((t for t in approx if year(t) < recent_since),
                   key=lambda t: (is_offscope(t), t["last_updated"] or "0000", t["country"]))
    return recent, stale, blank


def is_offscope(t):
    return all(f.strip().lower() not in LNG_FUELS for f in t["fuel"].split(";")) if t["fuel"] else False


def _neg(datestr):
    """Sort key that puts the newest date first without reversing the whole tuple."""
    return tuple(-ord(c) for c in (datestr or ""))


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, recent, stale, blank, terminals, n_unit_rows, recent_since):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    YELLOW = PatternFill("solid", fgColor="FFEB9C")
    RED = PatternFill("solid", fgColor="FFC7CE")
    BASE = Font(size=10)
    BOLD = Font(size=10, bold=True)
    GREY_FONT = Font(size=10, color="909090", italic=True)
    LINK = Font(size=10, color="0563C1", underline="single")
    TOP = Alignment(vertical="top")
    WRAP = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    col_no = {name: i + 1 for i, name in enumerate(COLS)}

    def sheet(wb, title, rows):
        ws = wb.create_sheet(title)
        ws.append(COLS)
        for c in range(1, len(COLS) + 1):
            cc = ws.cell(1, c)
            cc.fill, cc.font, cc.border = HDR_FILL, HDR_FONT, BORDER
            cc.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r[c] for c in COLS])
            rn = ws.max_row
            for c in range(1, len(COLS) + 1):
                cc = ws.cell(rn, c)
                cc.font, cc.alignment, cc.border = BASE, TOP, BORDER
            for key in ("project_url", "wiki", "location_ref"):
                val = r[key]
                if isinstance(val, str) and val.startswith("http") and " " not in val:
                    cc = ws.cell(rn, col_no[key])
                    cc.hyperlink, cc.font = val, LINK
                ws.cell(rn, col_no[key]).alignment = WRAP
            # coarse coordinates and missing ones are the priority targets
            if not r["latitude"] or not r["longitude"]:
                for key in ("latitude", "longitude"):
                    ws.cell(rn, col_no[key]).fill = RED
            elif r["coord_dp"] != "" and r["coord_dp"] <= 2:
                ws.cell(rn, col_no["coord_dp"]).fill = YELLOW
            if r["accuracy_mixed"]:
                ws.cell(rn, col_no["accuracy_mixed"]).fill = RED
            if is_offscope(r):
                for c in range(1, len(COLS) + 1):
                    ws.cell(rn, c).font = GREY_FONT
        for i, name in enumerate(COLS, 1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS[name]
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"
        return ws

    wb = Workbook()
    wb.remove(wb.active)

    sw = wb.create_sheet("summary")
    n_approx = sum(1 for t in terminals if t["accuracy"] == "approximate")
    n_exact = sum(1 for t in terminals if t["accuracy"] == "exact")
    lines = [
        ("LNG terminal location-accuracy audit", ""),
        ("", ""),
        (f"unit rows in export", n_unit_rows),
        ("distinct terminals", len(terminals)),
        ("  accuracy = exact", n_exact),
        ("  accuracy = approximate", n_approx),
        ("  accuracy blank", len(blank)),
        ("", ""),
        (f"approximate AND LastUpdated >= {recent_since}  -> sheet 'recent_approximate'", len(recent)),
        (f"    of which out-of-LNG-scope fuel (oil/NGL legacy records)", sum(1 for t in recent if is_offscope(t))),
        (f"approximate AND LastUpdated <  {recent_since}  -> sheet 'stale_approximate'", len(stale)),
        (f"    of which out-of-LNG-scope fuel (oil/NGL legacy records)", sum(1 for t in stale if is_offscope(t))),
        ("accuracy blank (neither value set)  -> sheet 'accuracy_blank'", len(blank)),
        ("", ""),
        ("one row per terminal; coordinates/Location/Accuracy are project-level fields,", ""),
        ("LastUpdated is per-unit-row (newest shown; 'MIXED' flags disagreement).", ""),
        ("coord_dp = decimal places on the coarser of lat/lon (<=2 highlighted).", ""),
        ("red latitude/longitude = coordinate missing entirely.", ""),
        ("grey italic rows = non-LNG fuel (crude/NGL deepwater ports), sorted last.", ""),
    ]
    for label, val in lines:
        sw.append([label, val])
    for row in sw.iter_rows():
        for cc in row:
            cc.font = BASE
    sw.cell(1, 1).font = Font(size=12, bold=True)
    for rn in (3, 4, 9, 11, 13):
        sw.cell(rn, 2).font = BOLD
    sw.column_dimensions["A"].width = 74
    sw.column_dimensions["B"].width = 10

    sheet(wb, "recent_approximate", recent)
    sheet(wb, "stale_approximate", stale)
    sheet(wb, "accuracy_blank", blank)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", default="gem_export.csv", help="fresh GEM export (default gem_export.csv)")
    ap.add_argument("--recent-since", type=int, default=2026,
                    help="LastUpdated year at/after which a record counts as recent (default 2026)")
    ap.add_argument("--out-dir", default="../batches/deliverables",
                    help="output directory (default ../batches/deliverables)")
    ap.add_argument("--stem", default=None, help="override output filename stem")
    args = ap.parse_args()

    n_unit_rows, terminals = rollup(args.csv)
    recent, stale, blank = split(terminals, args.recent_since)

    stamp = subprocess.run(["date", "+%Y%m%d_%H%M_ET"], env={**os.environ, "TZ": "America/New_York"},
                           capture_output=True, text=True, check=True).stdout.strip()
    stem = args.stem or f"location_accuracy_audit_{stamp}"
    os.makedirs(args.out_dir, exist_ok=True)
    base = os.path.join(args.out_dir, stem)

    write_xlsx(base + ".xlsx", recent, stale, blank, terminals, n_unit_rows, args.recent_since)
    write_csv(base + "_recent_approximate.csv", recent)
    write_csv(base + "_stale_approximate.csv", stale)

    print(f"{n_unit_rows} unit rows -> {len(terminals)} terminals")
    print(f"  approximate, LastUpdated >= {args.recent_since}: {len(recent)}")
    print(f"  approximate, LastUpdated <  {args.recent_since}: {len(stale)}")
    print(f"  accuracy blank: {len(blank)}")
    print(f"  -> {base}.xlsx (+ 2 csv)")


if __name__ == "__main__":
    main()
