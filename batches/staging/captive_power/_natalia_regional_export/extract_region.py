#!/usr/bin/env python3
"""Extract the captive-True (CaptiveGasPower=True) roster + all staged detail for
one regional captive workbook, for reformatting into Natalia's Qualifying/Excluded
threshold layout.

Usage: python3 extract_region.py <workbook.xlsx> <region-key> [<outdir>]

Writes <outdir>/<region-key>_detail.json:
  { TerminalID: {region, name, country, statuses[], cats[], hw[], refs[], notes[]} }

Terminals whose ID appears in exclude_ids.json (Natalia's original file + the
already-delivered missing_terminals companion) are skipped.
"""
import json, os, sys
import openpyxl

wb_path, region = sys.argv[1], sys.argv[2]
outdir = sys.argv[3] if len(sys.argv) > 3 else os.path.dirname(os.path.abspath(__file__))

exclude_path = os.path.join(outdir, 'exclude_ids.json')
exclude = set(json.load(open(exclude_path))) if os.path.exists(exclude_path) else set()

wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)

# YES roster + notes/refs from updates_summary (staged CaptiveGasPower=True rows)
us = wb['updates_summary']
rows = us.iter_rows(values_only=True)
h = {c: i for i, c in enumerate(next(rows))}
det = {}
for r in rows:
    if r[h['field_name']] == 'CaptiveGasPower' and str(r[h['new_value']]) == 'True':
        tid = r[h['terminal_id']]
        if tid in exclude:
            continue
        d = det.setdefault(tid, {'region': region, 'name': r[h['terminal_name']],
                                 'country': '', 'statuses': [], 'cats': [], 'hw': [],
                                 'refs': [], 'notes': []})
        note = r[h.get('source_notes', h.get('notes', 0))]
        if note and note not in d['notes']:
            d['notes'].append(str(note))
        for u in str(r[h['ref_urls']] or '').split('\n'):
            u = u.strip()
            if u and u not in d['refs']:
                d['refs'].append(u)

# category / hardware / status / country / ref from updates_in_database_format
db = wb['updates_in_database_format']
rows = db.iter_rows(values_only=True)
h2 = {c: i for i, c in enumerate(next(rows))}
for r in rows:
    tid = r[h2['TerminalID']]
    if tid not in det:
        continue
    d = det[tid]
    if r[h2.get('Country/Area')] and not d['country']:
        d['country'] = str(r[h2['Country/Area']])
    for key, field in (('statuses', 'Status'), ('cats', 'captive_category'),
                       ('hw', 'hardware_summary')):
        v = r[h2.get(field)] if h2.get(field) is not None else None
        if v and str(v) not in d[key]:
            d[key].append(str(v))
    ref = r[h2.get('CaptiveGasPower [ref]')] if 'CaptiveGasPower [ref]' in h2 else None
    if ref:
        for u in str(ref).split('\n'):
            u = u.strip()
            if u and u not in d['refs']:
                d['refs'].append(u)

out = os.path.join(outdir, f'{region}_detail.json')
json.dump(det, open(out, 'w'), indent=1, ensure_ascii=False)
print(f'{region}: {len(det)} captive-True terminals -> {out}')
for tid, d in sorted(det.items(), key=lambda kv: (kv[1]['country'], kv[1]['name'])):
    print(f"  {tid}  {d['country']:<22} {d['name']}  [{'; '.join(d['statuses'])}]")
