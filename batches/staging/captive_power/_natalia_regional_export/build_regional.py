#!/usr/bin/env python3
"""Build one region's Natalia-format workbook from <region>_classified.json +
<region>_detail.json (refs come from the detail file, keyed by terminal_id, so
classification agents never hand-copy URLs).

Usage: python3 build_regional.py <region-key> [<outdir>]
Region keys: asia, middle-east-gulf, africa, oceania.
Writes "Captive PPs_<Label>_08.02.2026.xlsx" into <outdir>.
"""
import json, os, sys
import openpyxl
from openpyxl.styles import Font, Alignment

REGION_LABELS = {'asia': ('Asia', 'Asia'),
                 'middle-east-gulf': ('Middle East-Gulf', 'ME-Gulf'),
                 'africa': ('Africa', 'Africa'),
                 'oceania': ('Oceania', 'Oceania')}

region = sys.argv[1]
outdir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(__file__))
file_label, tab_label = REGION_LABELS[region]

detail = json.load(open(os.path.join(outdir, f'{region}_detail.json')))
cls = json.load(open(os.path.join(outdir, f'{region}_classified.json')))

# every terminal classified exactly once, and none invented
claimed = [r['terminal_id'] for r in cls['qualifying']] + [r['terminal_id'] for r in cls['excluded']]
missing = set(detail) - set(claimed)
extra = set(claimed) - set(detail)
dupes = {t for t in claimed if claimed.count(t) > 1}
if missing or extra or dupes:
    sys.exit(f'ROSTER MISMATCH {region}: missing={missing} extra={extra} dupes={dupes}')

def refs(tid):
    return '\n'.join(detail[tid]['refs'])

Q_HDR = ['Terminal', 'Terminal ID', 'Country', 'Status', 'Hardware Type',
         'Individual Unit MW', 'Aggregate/Total MW', 'Qualifying Basis', 'References']
X_HDR = ['Terminal', 'Terminal ID', 'Country', 'Reason Excluded', 'MW Figure (if any)', 'References']

wb = openpyxl.Workbook()
wb.remove(wb.active)
FONT, BOLD = Font(name='Calibri', size=10), Font(name='Calibri', size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')

def add_sheet(title, hdr, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(hdr)
    for c in ws[1]:
        c.font = BOLD
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = FONT, WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

q_rows = [[r['terminal'], r['terminal_id'], r['country'], r['status'], r['hardware_type'],
           r['individual_mw'], r['aggregate_mw'], r['basis'], refs(r['terminal_id'])]
          for r in sorted(cls['qualifying'], key=lambda r: (r['country'], r['terminal']))]
x_rows = [[r['terminal'], r['terminal_id'], r['country'], r['reason'], r['mw_figure'],
           refs(r['terminal_id'])]
          for r in sorted(cls['excluded'], key=lambda r: (r['country'], r['terminal']))]

add_sheet(f'{tab_label} - Qualifying (>=50MW)', Q_HDR, q_rows, [30, 16, 18, 26, 55, 34, 30, 42, 55])
add_sheet(f'{tab_label} - Excluded', X_HDR, x_rows, [34, 16, 18, 75, 34, 55])

out = os.path.join(outdir, f'Captive PPs_{file_label}_08.02.2026.xlsx')
wb.save(out)
print(f'{region}: {len(q_rows)} qualifying + {len(x_rows)} excluded -> {out}')
